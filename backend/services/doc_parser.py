"""
Document Ingestion & Text Extraction Service
Extracts text from PDF, DOCX, and TXT files with sanitization.
"""

import io
import re
from typing import Tuple


class DocumentParser:
    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str]:
        """
        Extract clean text and determine document format.
        Returns: (extracted_text, file_type)
        """
        ext = filename.lower().split(".")[-1] if "." in filename else "txt"

        if ext == "pdf":
            text = DocumentParser._parse_pdf(file_bytes)
            file_type = "PDF"
        elif ext in ["docx", "doc"]:
            text = DocumentParser._parse_docx(file_bytes)
            file_type = "DOCX"
        else:
            text = DocumentParser._parse_plain_text(file_bytes)
            file_type = "TXT"

        cleaned_text = DocumentParser.clean_and_sanitize_text(text)
        return cleaned_text, file_type

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages)
        except Exception as e:
            # Fallback to UTF-8 / ASCII decode
            try:
                return file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                return f"Error parsing PDF content: {str(e)}"

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        except Exception as e:
            try:
                return file_bytes.decode("utf-8", errors="ignore")
            except Exception:
                return f"Error parsing DOCX content: {str(e)}"

    @staticmethod
    def _parse_plain_text(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1")
            except Exception:
                return file_bytes.decode("utf-8", errors="ignore")

    @staticmethod
    def clean_and_sanitize_text(raw_text: str) -> str:
        if not raw_text:
            return ""
        # Remove null bytes, non-printable control characters (except newline, tab)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', ' ', raw_text)
        # Normalize multiple spaces and multiple blank lines
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n\s*\n\s*\n+', '\n\n', text)
        return text.strip()


document_parser = DocumentParser()
