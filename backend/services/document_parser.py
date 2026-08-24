import os
import re
import io
from typing import Tuple
from werkzeug.utils import secure_filename

# Optional PDF / DOCX parsers with graceful fallbacks
try:
    from pypdf import PdfReader
except ImportError:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        PdfReader = None

try:
    import docx
except ImportError:
    docx = None


class DocumentParserService:
    ALLOWED_EXTENSIONS = {'pdf', 'docx', 'doc', 'txt'}

    @classmethod
    def is_allowed_file(cls, filename: str) -> bool:
        return '.' in filename and filename.rsplit('.', 1)[1].lower() in cls.ALLOWED_EXTENSIONS

    @classmethod
    def clean_text(cls, text: str) -> str:
        if not text:
            return ""
        # Remove ASCII control characters and non-printable noise
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', ' ', text)
        # Normalize multiple spaces, tabs, and excess newlines
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    @classmethod
    def parse_pdf(cls, file_stream: bytes) -> str:
        if PdfReader is None:
            raise RuntimeError("PDF parsing library (pypdf) is not installed.")
        try:
            reader = PdfReader(io.BytesIO(file_stream))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text)
            return cls.clean_text("\n\n".join(extracted_pages))
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    @classmethod
    def parse_docx(cls, file_stream: bytes) -> str:
        if docx is None:
            raise RuntimeError("DOCX parsing library (python-docx) is not installed.")
        try:
            doc = docx.Document(io.BytesIO(file_stream))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
            return cls.clean_text("\n\n".join(full_text))
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def parse_txt(cls, file_stream: bytes) -> str:
        for encoding in ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']:
            try:
                text = file_stream.decode(encoding)
                return cls.clean_text(text)
            except UnicodeDecodeError:
                continue
        return cls.clean_text(file_stream.decode('utf-8', errors='ignore'))

    @classmethod
    def parse_uploaded_file(cls, file_storage, upload_folder: str = None) -> Tuple[str, str]:
        """
        Parses a Werkzeug FileStorage object.
        Returns: (extracted_text, safe_filename)
        """
        raw_filename = file_storage.filename or "uploaded_tender.txt"
        safe_name = secure_filename(raw_filename)
        ext = safe_name.rsplit('.', 1)[1].lower() if '.' in safe_name else 'txt'

        file_bytes = file_storage.read()

        if upload_folder and os.path.exists(upload_folder):
            save_path = os.path.join(upload_folder, safe_name)
            with open(save_path, 'wb') as f:
                f.write(file_bytes)

        if ext == 'pdf':
            text = cls.parse_pdf(file_bytes)
        elif ext in ['docx', 'doc']:
            text = cls.parse_docx(file_bytes)
        else:
            text = cls.parse_txt(file_bytes)

        return text, safe_name
